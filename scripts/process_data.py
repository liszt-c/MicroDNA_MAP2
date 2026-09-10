"""
scripts/process_data.py - 数据预处理: 从 Excel 坐标表批量提取序列, 写入合并 FASTA
"""
import argparse
import sys
from collections import Counter
from pathlib import Path

import openpyxl

sys.path.append(str(Path(__file__).resolve().parent.parent))

from config import PROCESSED_DATA_DIR, HG19_FA, SAMTOOLS, SEQUENCE_LENGTH
from src.utils import setup_logger, ensure_faidx, extract_regions, read_fai

logger = setup_logger('process_data')

COL_ALIASES = {
    'chrom': ('chrom', 'chr', 'chromosome', 'contig', 'seqname', 'seq_name', 'reference'),
    'start': ('start', 'begin', 'left', 'st', 'start_pos', 'startpos'),
    'end':   ('end', 'stop', 'right', 'en', 'end_pos', 'endpos'),
}

LENGTH_BINS = [(0, 200), (200, 400), (400, 600), (600, 800), (800, 1000), (1000, None)]


# ---------------------------------------------------------------------- #
def detect_columns(header_values, header_row: int):
    """从表头行自动推断 chrom/start/end 所在列号 (1-based)"""
    names = {}
    for i, v in enumerate(header_values, start=1):
        if v is None:
            continue
        names[i] = str(v).strip().lower().replace(' ', '_').replace('.', '_')

    result = {}
    for key, aliases in COL_ALIASES.items():              # 先精确匹配
        for col, name in names.items():
            if name in aliases:
                result[key] = col
                break
    for key, aliases in COL_ALIASES.items():              # 再子串匹配
        if key in result:
            continue
        for col, name in names.items():
            if col in result.values():
                continue
            if any(a in name for a in aliases):
                result[key] = col
                break

    if len(result) < 3:
        logger.error(f"Auto-detection failed (row {header_row}). Detected: {result}. "
                     f"Headers: {names}. Please specify --chrom-col/--start-col/--end-col.")
        return None
    logger.info(f"Detected columns (row {header_row}): "
                f"chrom=col{result['chrom']}, start=col{result['start']}, end=col{result['end']}")
    return result


def make_chrom_resolver(fai_lengths: dict):
    """依据 .fai 索引解析染色体名 (自动尝试加/去 'chr' 前缀、大小写)"""
    known = set(fai_lengths)
    cache = {}

    def resolve(chrom: str):
        chrom = str(chrom).strip()
        if not chrom or chrom.lower() in ('nan', 'none'):
            return None
        if chrom in cache:
            return cache[chrom]
        bare = chrom[3:] if chrom.lower().startswith('chr') else chrom
        candidates = [chrom, f"chr{bare}", bare, chrom.upper(), chrom.lower(),
                      f"chr{bare.upper()}", f"chr{bare.lower()}"]
        hit = next((c for c in dict.fromkeys(candidates) if c in known), None)
        cache[chrom] = hit
        return hit

    return resolve


def adjust_region(start: int, end: int, mode: str, target_len: int,
                  min_span: int, chrom_len: int):
    """
    按策略调整区域坐标 (1-based 闭区间)
    :return: (new_start, new_end) 或 None (表示丢弃该记录)
    """
    span = end - start
    if span <= 0:
        return None

    if mode == 'raw':
        s, e = start, end

    elif mode == 'expand':
        if span >= target_len:
            s, e = start, end                      # 长片段原样保留, 编码时截取正中间
        else:
            # 复现 count_V5: half = span//2; fill = target/2 - half; 两侧各扩 fill
            half = span // 2
            fill = target_len / 2.0 - half
            s = int(start - fill)
            e = int(end + fill)
            if e - s == target_len + 1:            # 奇数跨度时多出的 1bp
                e -= 1

    elif mode == 'middle':
        if span <= min_span:
            return None                            # 复现 cout_other_v3: 跨度不足则跳过
        half = span // 2
        s = start + half - target_len // 2
        e = s + target_len

    else:
        raise ValueError(f"Unknown mode: {mode}")

    # 裁剪到染色体范围内
    s = max(1, min(s, chrom_len))
    e = max(s, min(e, chrom_len))
    return s, e


def read_regions(xlsx: Path, sheet, first_row, last_row, header_row,
                 chrom_col, start_col, end_col, mode, target_len, min_span,
                 ffill_chrom, resolve_chrom, fai_lengths, max_records):
    """读取 Excel 并生成待提取区域列表"""
    wb = openpyxl.load_workbook(str(xlsx), read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.active
        logger.info(f"Reading sheet '{ws.title}' from {xlsx.name}")

        if not all([chrom_col, start_col, end_col]):
            hdr = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
            detected = detect_columns(hdr, header_row)
            if detected is None:
                return [], Counter()
            chrom_col = chrom_col or detected['chrom']
            start_col = start_col or detected['start']
            end_col = end_col or detected['end']

        max_col = max(chrom_col, start_col, end_col)
        if last_row is None or last_row <= 0:
            last_row = ws.max_row

        regions = []
        seen = set()
        stats = Counter()
        last_chrom = None

        for row in ws.iter_rows(min_row=first_row, max_row=last_row,
                                max_col=max_col, values_only=True):
            if max_records and len(regions) >= max_records:
                logger.info(f"Reached --max-records {max_records}, stop reading.")
                break

            raw_chrom = row[chrom_col - 1] if chrom_col - 1 < len(row) else None
            raw_start = row[start_col - 1] if start_col - 1 < len(row) else None
            raw_end = row[end_col - 1] if end_col - 1 < len(row) else None

            # 染色体列可能是合并单元格 -> 沿用上一有效值
            if raw_chrom is None or str(raw_chrom).strip() == '':
                if ffill_chrom and last_chrom is not None:
                    raw_chrom = last_chrom
                    stats['chrom_ffilled'] += 1
                else:
                    stats['skipped_no_chrom'] += 1
                    continue
            last_chrom = raw_chrom

            try:
                start = int(float(raw_start))
                end = int(float(raw_end))
            except (TypeError, ValueError):
                stats['skipped_bad_coord'] += 1
                continue

            chrom = resolve_chrom(raw_chrom)
            if chrom is None:
                stats['skipped_unknown_chrom'] += 1
                continue
            if chrom != str(raw_chrom).strip():
                stats['chrom_renamed'] += 1

            chrom_len = fai_lengths.get(chrom, 0)
            if chrom_len <= 0:
                stats['skipped_unknown_chrom'] += 1
                continue

            adjusted = adjust_region(start, end, mode, target_len, min_span, chrom_len)
            if adjusted is None:
                stats['skipped_by_mode'] += 1
                continue
            s, e = adjusted

            orig_len = end - start
            for lo, hi in LENGTH_BINS:
                if orig_len >= lo and (hi is None or orig_len < hi):
                    stats[f'len_{lo}_{"inf" if hi is None else hi}'] += 1
                    break

            key = (chrom, s, e)
            if key in seen:
                stats['duplicated'] += 1
                continue
            seen.add(key)
            regions.append((chrom, s, e, orig_len))

        return regions, stats
    finally:
        wb.close()


def print_length_stats(stats: Counter, total: int):
    if total <= 0:
        return
    logger.info("--- 原始区域长度分布 ---")
    for lo, hi in LENGTH_BINS:
        k = f'len_{lo}_{"inf" if hi is None else hi}'
        c = stats.get(k, 0)
        label = f"{lo}-{hi}bp" if hi else f">={lo}bp"
        logger.info(f"  {label:<12}: {c:>8d}  ({c / total:.2%})")


# ---------------------------------------------------------------------- #
def process_one(args, label_str: str, output_fa: Path):
    ref = Path(args.ref)
    ensure_faidx(ref, SAMTOOLS)
    fai = read_fai(ref)
    if not fai['lengths']:
        raise RuntimeError(f"Failed to read .fai index for {ref}")
    logger.info(f"Reference: {ref} ({len(fai['lengths'])} contigs)")

    resolve_chrom = make_chrom_resolver(fai['lengths'])
    regions, stats = read_regions(
        xlsx=Path(args.input), sheet=args.sheet,
        first_row=args.first_row, last_row=args.last_row, header_row=args.header_row,
        chrom_col=args.chrom_col, start_col=args.start_col, end_col=args.end_col,
        mode=args.mode, target_len=args.target_len, min_span=args.min_span,
        ffill_chrom=args.ffill_chrom, resolve_chrom=resolve_chrom,
        fai_lengths=fai['lengths'], max_records=args.max_records)

    total_span = sum(e - s for _, s, e, _ in regions)
    logger.info(f"Parsed {len(regions)} unique regions "
                f"(total {total_span / 1e6:.2f} Mb); skipped: "
                f"{ {k: v for k, v in stats.items() if k.startswith('skipped')} }")
    print_length_stats(stats, sum(v for k, v in stats.items() if k.startswith('len_')))

    if not regions:
        logger.error("No regions to extract. Nothing written.")
        return 0

    keys = [(c, s, e) for c, s, e, _ in regions]
    seq_map, skipped = extract_regions(ref, keys, SAMTOOLS)
    for region, reason in skipped:
        logger.warning(f"Skipped {region[0]}:{region[1]}-{region[2]} -> {reason}")

    write_mode = 'a' if args.append else 'w'
    if not args.append and output_fa.exists():
        logger.info(f"Overwriting existing {output_fa}")

    sample_id = args.sample_id or Path(args.input).stem
    written = 0
    with open(output_fa, write_mode, encoding='utf-8') as f:
        for idx, (chrom, s, e, orig_len) in enumerate(regions):
            seq = seq_map.get((chrom, s, e))
            if not seq:
                continue
            seq = seq.upper()
            header = f">{sample_id}_{idx}|{chrom}:{s}-{e}|{label_str}"
            f.write(header + "\n")
            for i in range(0, len(seq), 70):
                f.write(seq[i:i + 70] + "\n")
            written += 1

    logger.info(f"[OK] Wrote {written} '{label_str}' sequences -> {output_fa}")
    return written


def main():
    p = argparse.ArgumentParser(
        description="Extract sequences from an Excel coordinate table into a merged FASTA.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    p.add_argument('--input', required=True, help="Excel 坐标表 (.xlsx)")
    p.add_argument('--label', choices=['ecc', 'other'], default='ecc',
                   help="写入 header 的标签: ecc -> eccDNA, other -> otherDNA")
    p.add_argument('--output', default=None,
                   help="输出 FASTA (默认 data/processed/{eccDNA,otherDNA}.fa)")
    p.add_argument('--ref', default=str(HG19_FA), help="参考基因组 FASTA (需可被 samtools faidx 索引)")
    p.add_argument('--sample-id', default=None, help="Header 中的样本 ID (默认取 Excel 文件名)")
    p.add_argument('--mode', choices=['expand', 'middle', 'raw'], default='expand',
                   help="区域长度调整策略: expand=向两侧扩至目标长度(真实侧翼); "
                        "middle=取跨度>min-span区域的正中间; raw=原样")
    p.add_argument('--target-len', type=int, default=SEQUENCE_LENGTH, help="目标长度 (bp)")
    p.add_argument('--min-span', type=int, default=600, help="middle 模式的最小跨度阈值 (bp)")

    p.add_argument('--sheet', default=None, help="工作表名 (默认活动表)")
    p.add_argument('--header-row', type=int, default=1, help="表头所在行 (1-based)")
    p.add_argument('--first-row', type=int, default=2, help="数据起始行 (1-based)")
    p.add_argument('--last-row', type=int, default=None, help="数据终止行 (默认到末尾)")
    p.add_argument('--chrom-col', type=int, default=None, help="染色体列号 (1-based, 默认自动检测)")
    p.add_argument('--start-col', type=int, default=None, help="起始坐标列号 (1-based)")
    p.add_argument('--end-col', type=int, default=None, help="终止坐标列号 (1-based)")
    p.add_argument('--ffill-chrom', action='store_true',
                   help="染色体单元格为空时沿用上一行的值 (处理合并单元格导出)")

    p.add_argument('--append', action='store_true', help="追加写入而非覆盖")
    p.add_argument('--max-records', type=int, default=0, help="仅处理前 N 条 (0=全部, 用于小规模测试)")
    args = p.parse_args()

    label_str = 'eccDNA' if args.label == 'ecc' else 'otherDNA'
    if args.output:
        output_fa = Path(args.output)
    else:
        output_fa = PROCESSED_DATA_DIR / f"{label_str}.fa"
    output_fa.parent.mkdir(parents=True, exist_ok=True)

    n = process_one(args, label_str, output_fa)
    if n == 0:
        sys.exit(1)


if __name__ == '__main__':
    main()